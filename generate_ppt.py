from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


OUT = "visual_nav_station_overview.pptx"
BG = RGBColor(8, 18, 32)
PANEL = RGBColor(17, 35, 55)
CYAN = RGBColor(82, 215, 255)
GREEN = RGBColor(73, 226, 160)
WHITE = RGBColor(235, 243, 250)
MUTED = RGBColor(153, 176, 198)


def box(slide, x, y, w, h, fill=PANEL, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value; run.font.name = "Microsoft YaHei"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return tb


def bullets(slide, items, x, y, w, h, size=15):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = "• " + item; p.level = 0; p.space_after = Pt(7); p.font.name = "Microsoft YaHei"; p.font.size = Pt(size); p.font.color.rgb = WHITE
    return tb


def base_slide(prs, title, subtitle, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    box(slide, 0, 0, 13.333, 0.12, CYAN)
    text(slide, title, 0.55, 0.42, 9.8, 0.55, 27, WHITE, True)
    text(slide, subtitle, 0.58, 1.02, 11.4, 0.35, 11, MUTED)
    text(slide, f"VISUAL NAV STATION  /  {number:02d}", 10.5, 0.48, 2.25, 0.3, 10, CYAN, True, PP_ALIGN.RIGHT)
    return slide


prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

s = base_slide(prs, "视觉导航地面站", "面向双光视觉导航设备的 PC 上位机：真实 ROS 数据、远程配置与 GPS 对比", 1)
box(s, 0.58, 1.62, 5.65, 5.2, PANEL, RGBColor(31, 62, 88), True)
text(s, "当前系统定位", 0.88, 1.92, 4.6, 0.35, 18, CYAN, True)
bullets(s, ["PyQt5 多页面桌面架构，开发阶段以 main.py 快速迭代", "SSH 连接 Firefly 板卡，独立控制 roscore、传感器、MAVROS 与算法", "RGB / 红外双光模式二选一，适配白天与黑夜导航", "主页面实时显示 RGB、红外、IMU、ToF、飞控 GPS 与算法 GPS"], 0.9, 2.42, 4.9, 2.25, 14)
text(s, "数据链路", 0.88, 4.95, 4.6, 0.35, 18, GREEN, True)
text(s, "板卡 ROS 话题  →  RosBridge  →  主页面可视化\n\nGeoTIFF 底图  →  仿射坐标转换  →  FCU / GeoVINS 轨迹与误差", 0.9, 5.38, 4.95, 1.1, 14, WHITE)
box(s, 6.55, 1.62, 6.18, 5.2, RGBColor(10, 27, 44), RGBColor(31, 62, 88), True)
text(s, "MISSION CONTROL", 6.9, 1.92, 5.4, 0.35, 18, CYAN, True)
for i, (label, color) in enumerate([("RGB / IR 实时预览", CYAN), ("GPS 对比地图 + ROI 跟随", GREEN), ("ATTITUDE / FCU GPS / VEL 曲线", RGBColor(244, 201, 93)), ("启动 / 停止 / 日志诊断", RGBColor(255, 139, 104))]):
    yy = 2.52 + i * 0.78; box(s, 6.9, yy, 5.35, 0.56, RGBColor(18, 43, 64), color, True); text(s, label, 7.15, yy + 0.12, 4.8, 0.25, 14, WHITE, True)
text(s, "地图显示特性", 6.9, 5.8, 5.0, 0.3, 15, MUTED, True)
text(s, "固定窗口尺寸 · 等比例裁切 · + / − 缩放 · 跟随 GPS 开关 · 大地图窗口", 6.9, 6.18, 5.4, 0.42, 13, WHITE)

s = base_slide(prs, "已实现能力与下一阶段路线", "从工程样机走向可维护、可配置、可量产的视觉导航地面站", 2)
box(s, 0.58, 1.62, 5.95, 5.35, PANEL, RGBColor(31, 62, 88), True)
text(s, "第一版已实现", 0.9, 1.92, 4.8, 0.35, 18, CYAN, True)
bullets(s, ["传感器页面：单独启动 / 停止 / 话题测试，支持算法配置文件远程读写", "地图页面：地图资产上传板卡；客户端 GeoTIFF GPS 对比底图独立加载", "定位模式：串口 NMEA / ULBOX 与 MAVLink HIL_GPS 选项", "通讯诊断：SSH 内嵌终端、连接日志、断线重连入口", "真实数据优先：ROS Noetic、CompressedImage、IMU、GPS、ToF"], 0.92, 2.4, 5.05, 3.5, 14)
box(s, 6.78, 1.62, 5.95, 5.35, RGBColor(10, 27, 44), RGBColor(31, 62, 88), True)
text(s, "推荐后续功能", 7.1, 1.92, 4.8, 0.35, 18, GREEN, True)
bullets(s, ["设备配置中心：每台设备独立 YAML，统一保存 IP、话题、launch、算法版本", "进程与健康监控：节点状态、话题频率、延迟、CPU / 内存 / 温度", "地图工程管理：底图 CRS 校验、轨迹回放、误差曲线与 CSV 报告导出", "安全与量产：密钥认证、权限分级、配置备份 / 恢复、版本升级包", "软件封装：PyInstaller / deb 安装包、自动检查依赖与一键启动"], 7.12, 2.4, 5.0, 3.5, 14)
text(s, "建议优先级：健康监控 → 配置中心 → 轨迹回放与报告 → 量产封装", 7.12, 6.2, 5.0, 0.38, 13, CYAN, True)

prs.save(OUT)
print(OUT)
